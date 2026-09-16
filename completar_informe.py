from pathlib import Path
import json
import re
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import RGBColor, Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parent
doc=Document(r'C:/Users/joale/Downloads/Informe_Final_Pruebas_Skainet_v0.1 (2).docx')
paras=[Paragraph(e,doc) for e in doc.element.body.xpath('.//w:p')]
updates={
47:'La cobertura del criterio de aceptación aparece marcada con advertencia en 15 casos por trazabilidad afectada (ver 6.2 y 15).',
8:'Código local de Skainet_FI, configuración y pruebas Jest, reportes de cobertura y verificación de compilación/lint. Se conservan los resultados documentales del checklist consignados en la versión 0.1.',
10:'14 de septiembre de 2026',
12:'Web: React 19, JavaScript/JSX y Vite 8. API REST y worker: NestJS 11 y TypeScript 5. Persistencia: PostgreSQL con Prisma 5. App móvil: Flutter y Dart.',
14:'Revisión del informe documental, inspección estática del código y ejecución local de pruebas, compilación y lint. Evaluación del despliegue y análisis SonarQube pendientes.',
16:'Aprobación condicionada para entrega académica; no apto para producción hasta cerrar hallazgos críticos de seguridad y completar las validaciones pendientes.',
37:'El informe consolida los resultados documentales registrados en la versión 0.1 y la evaluación técnica de la copia local de Skainet. La revisión cubre arquitectura, autenticación, usuarios, inventario, producción, búsqueda y pruebas automatizadas. El antecedente documental contempla 22 requisitos funcionales, 22 casos de uso y 22 historias de usuario; sus cifras se mantienen separadas de los resultados de ejecución del software.',
38:'La versión 0.1 registra 21 de 22 RF con estado CUMPLE y uno con CUMPLE PARCIAL; los 22 CU figuran como CUMPLE. Las 22 HU aparecen como NO CUMPLE por ausencia de fichas formales. Estos indicadores describen la documentación evaluada en esa versión, no acreditan por sí solos que la implementación cumpla los requisitos.',
52:'239/239 pruebas aprobadas; 58/58 suites. Cobertura de líneas: 52,43 %.',
53:'Ejecución local con Jest. API, worker y frontend compilan. ESLint del frontend: 0 errores, 3 advertencias.',
54:'Conclusión ejecutiva: existe una base funcional y una suite reproducible, pero la aprobación sigue condicionada. La inspección identifica exposición potencial de datos de usuario, contraseñas sin hash, validaciones jerárquicas incompletas y riesgos de consistencia del inventario. SonarQube, las pruebas reales de despliegue y el cierre documental son requisitos pendientes para una aceptación final.',
67:'Requerimientos_Funcionales.docx (referenciado en la versión 0.1)',
69:'Antecedente documental de RF, RN y RNF recogido en el informe base.',
70:'Checklist_QA_Skaient.xlsx (referenciado en la versión 0.1)',
72:'Resultados documentales reproducidos de la versión 0.1: 22 RF, 22 CU y 22 HU.',
73:'Sistema_de_gestión_del_inventario_vs2.pptx (referenciado en la versión 0.1)',
75:'Contexto de negocio y catálogo preliminar descritos en el informe base.',
76:'https://github.com/Skainet-Code/Skainet_code_Completo.git',
77:'Copia local Skainet_FI; HEAD 9cc5ce95. Revisados Skainet-API, Skainet-Backend, Skainet-Frontend, manifiesto móvil y test/.',
78:'Inspección de la copia de trabajo local y del remoto configurado; no representa una auditoría del historial completo de Git.',
79:'Skainet-API/src/**/*.spec.ts; Skainet-API/pruebas/*.spec.ts; test/unitarias/*.spec.ts; test/integracion/*.spec.ts',
80:'58 archivos ejecutados por Skainet-API/jest.config.json; 239 resultados en informe-evidencias-jest.json.',
81:'Inventario completo por archivo en Anexo A. Las copias de pruebas no incluidas en esta configuración no se suman al total ejecutado.',
82:'Skainet-API/coverage/lcov.info y coverage-summary.json; reporte HTML en coverage/lcov-report/index.html.',
83:'Cobertura de código generada por Jest',
84:'473 de 902 líneas cubiertas; incluye archivos de producción sin ejecutar. No equivale a un análisis de SonarQube.',
85:'Alcance de la evidencia: los totales RF/CU/HU corresponden al antecedente documental de la versión 0.1. La evidencia técnica corresponde a la copia local revisada el 14/09/2026. La ejecución no incluye aceptación del sistema desplegado, disponibilidad, base de datos real aislada, Cypress ni Flutter.',
89:'Contraste de los resultados del checklist consignados en el informe base, manteniendo su procedencia documental.',
93:'Ejecución de Jest y revisión de cobertura; compilación de API, worker y frontend; lint del frontend. Inspección de las pruebas vinculadas con inventario, producción y pesaje.',
108:'Skainet separa una interfaz web, una API REST principal, un worker de automatización y telemetría y una aplicación móvil. La API incorpora usuarios, autenticación, inventario, órdenes, lotes, clientes, maquinaria, alertas, auditoría, búsqueda, estadísticas, recuperación y registros de producción. Las tecnologías se verificaron en los manifiestos y archivos de entrada del proyecto.',
112:'Frontend web',113:'React 19; React Router 7; Vite 8; JavaScript/JSX; Lucide React. package.json y src/App.jsx.',
114:'Paneles de usuarios e inventario, búsqueda global, pedidos, PhaseTimerWidget y TripleWeightModal. Build aprobado; Cypress presente, no ejecutado.',
115:'API y worker',116:'NestJS 11, TypeScript 5, Passport/JWT. Skainet-API/src y Skainet-Backend/src.',
117:'API principal en puerto 3000 por defecto; worker con módulos automation y telemetry. Ambos compilan; la suite de 239 casos corresponde a la API.',
118:'Persistencia',119:'PostgreSQL y Prisma 5. Esquema en Skainet-API/prisma/schema.prisma; conexión mediante DATABASE_URL.',
120:'Modelos User, AuditLog, Material, KardexMovement, PhaseTimeLog y TripleWeightLog, entre otros. No se ejecutaron migraciones contra una BD real en esta revisión.',
121:'Contrato de API y app móvil',122:'Swagger/OpenAPI en /api mediante @nestjs/swagger. Flutter/Dart con Provider, Dio, HTTP y almacenamiento seguro.',
123:'Documentación de endpoints y Bearer Auth configurada en src/main.ts. App móvil presente; compilación y pruebas Flutter fuera de la ejecución realizada.',
148:'Verificación parcial: UsersService contiene validaciones de roles y autodesactivación, pero RN-001/RN-002/RN-011/RN-014 no pueden darse por satisfechas. createUser permite al Super Administrador crear joyeros; updateUser y deactivateUser no bloquean explícitamente al Administrador frente a un usuario Super Administrador. Además, JwtStrategy retorna un usuario con id mientras los controladores toman req.user.sub y recurren a datos del cuerpo o a SISTEMA, lo que afecta la identidad del actor y la comprobación de autodesactivación. RN-008 se verifica al iniciar sesión, pero JwtStrategy no revisa accountStatus. El bloqueo de intentos existe en una rama alternativa de AuthService que no se alcanza cuando está disponible UsersService.validatePassword (TEC-003 y TEC-004).',
226:'Ejecución del 14/09/2026: 239 pruebas aprobadas, cero fallidas y 58 suites aprobadas (14,57 s en la ejecución con reporte JSON). La cobertura obtenida en la ejecución con instrumentación fue: sentencias 51,16 %, ramas 53,73 %, funciones 50,93 % y líneas 52,43 %. Los grupos siguientes son excluyentes y suman 239; el inventario detallado está en Anexo A.',
231:'Usuarios, usuarios-joyeros, autenticación; pruebas users y auth de src/pruebas.',232:'93',
234:'test/unitarias/inventario.unit.spec.ts y test/integracion/inventario.integration.spec.ts',235:'10',
237:'Sin prueba directa de ProductionLogsService.logPhaseTime en la suite ejecutada.',238:'0 directos',
240:'test/{unitarias,integracion}/produccion-mermas.*.spec.ts; cubren órdenes/mermas, no logTripleWeight.',241:'8 relacionados',
243:'test/unitarias/busqueda.unit.spec.ts y test/integracion/busqueda.integration.spec.ts',244:'7',
247:'239',
254:'Pruebas de altas, modificación, desactivación y permisos con dependencias simuladas; users.service.ts: 86,01 % de líneas.',255:'Parcial; faltan accesos HTTP reales y cierre de TEC-004.',
257:'17 casos de autenticación entre suite unitaria, suite denominada integración y pruebas AuthService. auth.service.ts: 87,75 % de líneas.',258:'Parcial; mocks no acreditan la ruta real de bloqueo ni firma JWT.',
260:'10 casos etiquetados como inventario. inventory.service.ts: 56,17 % de líneas. Persisten riesgos de atomicidad en stock/kardex.',261:'Parcial; requiere prueba transaccional y concurrente.',
263:'PhaseTimerWidget.jsx y logPhaseTime presentes. production-logs.service.ts tiene 0 % de cobertura.',264:'Implementado; sin evidencia automatizada directa.',
266:'TripleWeightModal.jsx y logTripleWeight presentes; 8 casos relacionados con producción/mermas de órdenes no validan el nuevo servicio.',267:'Parcial; logTripleWeight sin cobertura directa.',
269:'7 casos de búsqueda; search.service.ts: 100 % de líneas y 88,57 % de ramas.',270:'Servicio probado; permisos de extremo a extremo pendientes.',
272:'Build Vite aprobado; ESLint: 0 errores y 3 advertencias. Existen login.cy.js y navegacion.cy.js.',273:'Compilación/lint verificados; Cypress no ejecutado.',
275:'Hallazgos de inspección estática sobre la copia local. La severidad refleja impacto y prioridad de cierre; no son incidencias obtenidas de SonarQube ni explotación en un entorno desplegado. Las referencias permiten localizar la evidencia y reproducir la revisión.',
292:'Matriz integrada de hallazgos documentales y técnicos. Los hallazgos críticos condicionan cualquier liberación con datos reales.',
297:'TEC-001: lectura pública de usuario que puede incluir contraseña. TEC-002: contraseñas sin hash y secreto JWT predeterminado.',
298:'Eliminar datos sensibles de respuestas; proteger rutas; migrar contraseñas a hash robusto; exigir secreto JWT externo. Añadir pruebas HTTP negativas y verificar la configuración de despliegue.',
300:'TEC-003 a TEC-006: bloqueo de login, identidad/permisos, atomicidad del inventario y falta de cobertura de producción. Además, 22 HU pendientes y RF-005/RF-014 sin desarrollo documental según el antecedente.',
301:'Cerrar hallazgos con pruebas de integración reales y regresión; formalizar HU y requisitos de órdenes/alertas. La existencia de módulos en código no resuelve su falta de especificación.',
303:'TEC-007/TEC-008: configuración de pruebas del worker y precisión/tolerancia del pesaje. Persisten inconsistencias documentales RF-003/RF-004 y criterios RNF no medidos.',
304:'Corregir rutas de Jest del worker; aprobar tolerancia y unidad del pesaje; probar límites numéricos; unificar IDs y ejecutar mediciones RNF.',
306:'Fase 1 — Corrección inmediata: cerrar TEC-001 y TEC-002 antes de usar datos reales. Proteger consultas de usuarios, excluir password de las respuestas y reemplazar almacenamiento de contraseñas sin hash.',
307:'Fase 2 — Seguridad e integridad: corregir el flujo real de login, exigir JWT_SECRET, revisar caducidad/revocación, tomar el actor únicamente de la identidad autenticada y aplicar transacciones y control de concurrencia en stock/kardex.',
324:'Estado: configuración preparada en sonar-project.properties con projectKey skainet-api, fuentes Skainet-API/src e importación LCOV desde Skainet-API/coverage/lcov.info. No se ha ejecutado ni publicado un análisis en un servidor SonarQube; Quality Gate, deuda técnica, vulnerabilidades, duplicación y demás métricas de SonarQube no están evaluadas. La cobertura local de Jest (52,43 % de líneas; 53,73 % de ramas) no debe presentarse como un resultado del servidor. Para cerrar: generar cobertura con ejecutar_cobertura.bat, configurar el servidor/proyecto y token de análisis de forma privada, ejecutar SonarScanner desde la raíz y conservar el enlace y la captura del análisis procesado. LCOV transmite cobertura, no el conteo de pruebas.',
348:'Resultado RNF: RNF-001 no medido sobre endpoints reales; la prueba CP-178 crea 50 respuestas simuladas con duración aleatoria y no acredita carga del servidor. RNF-002 pendiente de verificar TLS del despliegue y de la conexión a BD. RNF-003 presenta incumplimiento en contraseñas: comparación y escritura sin hash en UsersService. RNF-004 no evaluado: requiere un periodo definido de monitoreo para calcular disponibilidad. RNF-005 no acreditado: el JWT expira a los 7 días y no demuestra expiración por inactividad. RNF-006 parcialmente implementado con toFixed(2), pero los pesos usan Float y el servicio de pesaje triple no tiene cobertura directa; faltan pruebas de precisión y límites. Las pruebas locales con mocks no sustituyen estas mediciones.',
355:'APROBADO. Vite 8.0.8: 1739 módulos transformados, build en 1,13 s y salida dist/. Comando: node node_modules/vite/bin/vite.js build.',
358:'APROBADO con advertencias. ESLint terminó con 0 errores y 3 advertencias en 13 archivos reportados. Cypress está excluido de esta configuración. Evidencia: informe-evidencias-lint.json.',
361:'APROBADO para compilación. API y worker: node node_modules/typescript/bin/tsc -p tsconfig.build.json, código de salida 0 en ambos.',
364:'APROBADO para la API: 58 suites y 239 pruebas aprobadas, 0 fallidas. Worker no incluido en este resultado; revisar TEC-007.',
367:'Pendiente de ejecución sobre BD aislada. No se realizaron migraciones ni seed contra el entorno de desarrollo.',
370:'NO ACREDITADO. Hay pruebas con mocks, pero TEC-004 exige corregir identidad y permisos y probar cada rol mediante HTTP.',
373:'PARCIAL. Inventario y órdenes tienen casos aprobados; cronometría y pesaje triple no tienen pruebas directas del servicio.',
379:'NO APROBADO para configuración segura: existe un secreto JWT de respaldo conocido y credenciales de demostración en código. Revisar paquete e historial antes de publicar; no se certifica ausencia de secretos reales.',
382:'Pendiente de cierre documental: según el informe base, 22 HU no cumplen y un RF cumple parcialmente.',
384:'La revisión técnica confirma que Skainet dispone de componentes implementados, compilación funcional y 239 pruebas aprobadas. Sin embargo, una cobertura de líneas del 52,43 % deja rutas relevantes sin verificar y la inspección detecta fallos de seguridad e integridad que requieren corrección. Los resultados documentales de la versión 0.1 se conservan como antecedente y deben actualizarse junto con los requisitos y las historias de usuario.',
385:'Resultado: aprobación condicionada para entrega académica, con impedimento para recomendar uso en producción mientras sigan abiertos TEC-001 y TEC-002. El cierre exige corregir permisos y autenticación, validar inventario transaccional, probar cronometría/pesaje triple, completar SonarQube y RNF y resolver las observaciones documentales. La aprobación de Jest y de compilación no implica aceptación integral del sistema.',
387:'La versión 0.1 deja registrada una comparación entre las hojas Checklist RF, Checklist CU, Checklist HU y Resumen. Se conserva ese antecedente documental y se incorpora la evidencia técnica de esta revisión en las secciones 5, 7, 8, 11, 12 y 13.',
389:'Según la comprobación registrada en la versión 0.1, los totales coinciden: RF 22, CU 22 y HU 22. No se sustituye ese conteo por el número de pruebas Jest, pues corresponden a unidades de evaluación diferentes.',
391:'Se mantienen las observaciones documentales del informe base en la sección 6.5 y Anexo B. Los nuevos hallazgos de implementación se registran como TEC-001 a TEC-008 para separar requisitos, código y resultados de pruebas.',
397:'Formalizar los umbrales de intentos fallidos y pesaje. El código compara el pesaje triple con 0,1, mientras su comentario menciona 0,05 g; el bloqueo de login de la rama alternativa tampoco acredita su aplicación en el flujo real.',
399:'Inventario de los 58 archivos efectivamente ejecutados por Jest; la columna de casos suma 239. Cada archivo terminó aprobado. Tipo indica ubicación de la suite y no certifica uso de servicios externos reales. Rutas relativas a la raíz Skainet_FI; nombres y resultados individuales disponibles en informe-evidencias-jest.json.',
}
for i,text in updates.items():
 p=paras[i]
 p.text=text
 for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)

# Replace the two placeholder findings and extend the same table.
findings=[
('TEC-001','Crítica','users.controller.ts: GET /users/:id sin guard; users.service.ts: findOne retorna parseUser(user), sin excluir password.','Posible lectura no autenticada de la contraseña almacenada y otros datos del usuario.','Proteger endpoint y aplicar selección/DTO de salida sin secretos; probar respuestas HTTP sin token y con cada rol.'),
('TEC-002','Crítica','UsersService escribe y compara password directamente. AuthModule y JwtStrategy admiten un secreto JWT fijo de respaldo.','Exposición de credenciales y riesgo de tokens falsificados si no se configura JWT_SECRET.','Hash robusto con sal por contraseña, migración de claves y secreto obligatorio gestionado fuera del código.'),
('TEC-003','Alta','AuthService.login retorna en la rama validatePassword; esa función no valida lockoutUntil ni incrementa intentos. La rama posterior contiene el bloqueo.','El flujo normal puede omitir bloqueo y auditoría pese a pruebas que pasan con mocks de otra rama.','Unificar el flujo y probar cinco intentos fallidos, bloqueo, expiración y auditoría con dependencias reales.'),
('TEC-004','Alta','JwtStrategy devuelve id; controladores leen sub y usan actorId del cuerpo o SISTEMA. Usuarios: falta bloqueo explícito sobre Super Administrador en update/deactivate.','Identidad del actor y autodesactivación inconsistentes; permisos jerárquicos incompletos.','Usar req.user.id, prohibir actor del cuerpo, validar matriz de permisos y estado activo en cada solicitud protegida.'),
('TEC-005','Alta','InventoryService.registerEntry/registerSalida actualizan stock y crean kardex en operaciones separadas; stock se calcula tras una lectura.','Un fallo intermedio o salidas concurrentes pueden producir stock y kardex inconsistentes.','Transacción, actualización atómica condicionada y pruebas de rollback/concurrencia en PostgreSQL aislado.'),
('TEC-006','Alta','production-logs.service.ts tiene 0 % de cobertura. CP-178 de rendimiento usa respuestas simuladas y tiempos aleatorios.','No se acreditan cronometría, pesaje triple ni capacidad real del servidor mediante esos resultados.','Pruebas directas del servicio, integración HTTP/BD y carga real con latencias medidas.'),
('TEC-007','Media','Skainet-Backend/jest.config.json apunta a backend/ en lugar de Skainet-Backend/.','Configuración no alineada con las rutas del worker; sus pruebas no forman parte de las 239 de la API.','Corregir rutas y ejecutar una suite independiente del worker, con evidencia separada.'),
('TEC-008','Media','logTripleWeight rechaza diferencias > 0,1; el comentario indica 0,05 g. Prisma modela pesos como Float y el servicio redondea a 2 decimales.','Ambigüedad de tolerancia y riesgo de representación numérica en registros de materiales preciosos.','Aprobar valor/unidad, alinear código y especificación; evaluar Decimal y probar fronteras, redondeo y valores no finitos.'),
]
ft=next(t for t in doc.tables if t.cell(0,0).text=='ID' and len(t.columns)==5)
for row in list(ft.rows)[1:]:ft._tbl.remove(row._tr)
for vals in findings:
 cells=ft.add_row().cells
 for c,v in zip(cells,vals):c.text=v

# Complete the area breakdown with the modules outside the five original rows.
area=next(t for t in doc.tables if t.cell(0,0).text=='Área')
row=area.add_row()
for c,v in zip(row.cells,['Otros módulos','Órdenes, lotes, maquinaria, recuperación, alertas, estadísticas, auditoría, clientes, configuración y otros.','121']):c.text=v
area.rows[-2]._tr.addprevious(row._tr)

# Complete annex with every executed suite, without counting duplicate folders.
report=json.loads((BASE/'informe-evidencias-jest.json').read_text(encoding='utf-8'))
at=next(t for t in doc.tables if t.cell(0,0).text=='#')
for c,v in zip(at.rows[0].cells,['#','Archivo','Casos','Ubicación y resultado']):c.text=v
for row in list(at.rows)[1:]:at._tbl.remove(row._tr)
for i,s in enumerate(sorted(report['testResults'],key=lambda x:x['name']),1):
 path=Path(s['name']).relative_to(BASE).as_posix()
 kind='Integración declarada' if '/integracion/' in path else 'Unitaria / componente'
 cells=at.add_row().cells
 for c,v in zip(cells,[str(i),path,str(len(s['assertionResults'])),kind+'; aprobado']):c.text=v
at.autofit=False
for col,width in zip(at.columns,[0.4,3.5,0.6,2.0]):col.width=Inches(width)
for row in at.rows:
 for c,width in zip(row.cells,[0.4,3.5,0.6,2.0]):c.width=Inches(width)

# Preserve original template, use black for completed text and readable table layout.
for p in doc.paragraphs:
 if re.match(r'^(\d+\.\s|\d+\.\d+\s|Anexo [AB]\.|Contenido$)',p.text) and len(p.text)<110 and p._p not in [paras[i]._p for i in range(19,36)]:
  p.paragraph_format.keep_with_next=True
  p.paragraph_format.keep_together=True
  for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
 if p.style and p.style.name.startswith('Heading'):
  for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
for s in doc.styles:
 if s.type==1 and (s.name.startswith('Heading') or s.name in ['Title','Subtitle']):s.font.color.rgb=RGBColor(0,0,0)
for table in doc.tables:
 for ri,row in enumerate(table.rows):
  trpr=row._tr.get_or_add_trPr()
  if not trpr.xpath('./w:cantSplit'):trpr.append(OxmlElement('w:cantSplit'))
  for h in row._tr.xpath('./w:trPr/w:trHeight'):h.getparent().remove(h)
  if ri==0:
   pr=row._tr.get_or_add_trPr()
   if not pr.xpath('./w:tblHeader'):pr.append(OxmlElement('w:tblHeader'))
  for c in row.cells:
   for p in c.paragraphs:
    p.paragraph_format.keep_with_next=False
    if table._tbl in [ft._tbl,at._tbl]:
     p.paragraph_format.space_after=Pt(3)
     p.paragraph_format.space_before=Pt(3)
     for r in p.runs:r.font.size=Pt(9.5)
    if ri>0:
     for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
   if table._tbl in [ft._tbl,at._tbl,area._tbl]:
    tcpr=c._tc.get_or_add_tcPr()
    margins=tcpr.find(qn('w:tcMar'))
    if margins is None:
     margins=OxmlElement('w:tcMar');tcpr.append(margins)
    for edge in ['top','left','bottom','right']:
     elem=OxmlElement('w:'+edge);elem.set(qn('w:w'),'75');elem.set(qn('w:type'),'dxa');margins.append(elem)

# Clear all unresolved red runs, asserting their placeholder text was replaced.
for p in doc.element.body.xpath('.//w:p'):
 text=''.join(p.xpath('.//w:t/text()'))
 assert '[Completar' not in text and '[Pendiente:' not in text,text
 for color in p.xpath('.//w:color'):
  if color.get(qn('w:val')) in ['C00000','FF0000']:color.set(qn('w:val'),'000000')
doc.core_properties.title='Informe final de pruebas y aseguramiento de calidad de Skainet'
for r in doc.paragraphs[0].runs:r.font.color.rgb=RGBColor(0,0,0)
doc.core_properties.subject='Evaluación técnica y resultados de pruebas del 14 de septiembre de 2026'
doc.save(BASE/'Informe_Final_Pruebas_Skainet_Completado.docx')
print('Created',BASE/'Informe_Final_Pruebas_Skainet_Completado.docx')
