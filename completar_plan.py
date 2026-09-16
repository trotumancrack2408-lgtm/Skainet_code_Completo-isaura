from pathlib import Path
import re
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

base=Path(__file__).resolve().parent
doc=Document(r'C:/Users/joale/OneDrive/Documents/Plan_Maestro_de_Pruebas_Skainet.docx')
ps=[Paragraph(e,doc) for e in doc.element.body.xpath('.//w:p')]
changes={
28:'Jose Alejandro Gualguan Rivera',30:'Jose Alejandro Gualguan Rivera',42:'Jose Alejandro Gualguan Rivera',
123:'Jose Alejandro Gualguan Rivera',126:'Jose Alejandro Gualguan Rivera',129:'Jose Alejandro Gualguan Rivera',132:'Jose Alejandro Gualguan Rivera',135:'Jose Alejandro Gualguan Rivera',138:'Jose Alejandro Gualguan Rivera',
6:'1.2 — actualización técnica',8:'14/09/2026',24:'14/09/2026; versión original: 17/07/2026',
37:'Instructor / Product Owner por confirmar',40:'Pendiente de aprobación',45:'Pendiente de aprobación',47:'Jose Alejandro Gualguan Rivera',50:'Pendiente de aprobación',
55:'ISO/IEC/IEEE 29119',56:'Referencia para organizar conceptos, proceso y documentación de pruebas. Se toman como orientación las partes 1:2022 y 2:2021 y la estructura documental de la serie. Este plan no declara una certificación de conformidad.',
57:'ISO/IEC 25010:2023',58:'Referencia para especificar y evaluar la calidad del producto. Se priorizan adecuación funcional, rendimiento, seguridad, fiabilidad y mantenibilidad con criterios medibles; citar esta norma no demuestra por sí solo que el producto cumpla sus características.',
60:'Este Plan Maestro define el alcance, la estrategia, los recursos, los datos y los criterios de cierre de las pruebas de Skainet. La actualización técnica del 14/09/2026 integra la implementación local y la evidencia de ejecución, distingue las pruebas realizadas de las planificadas y prioriza los riesgos detectados en el informe final.',
61:'Se conserva el Sprint QA-02 y su planificación original de 40 horas del 17 al 31 de julio de 2026 como línea base, sin afirmar que esas actividades se ejecutaron en esas fechas. Al corte técnico del 14/09/2026 se verificaron 239 pruebas Jest aprobadas en 58 suites, cobertura de líneas de 52,43 %, compilación de API/worker/frontend y lint web con 0 errores y 3 advertencias. Integración real, SonarQube, RNF del despliegue y aceptación de negocio continúan pendientes.',
64:'Validar criterios de aceptación: comprobar condiciones medibles aprobadas para cada requisito. Las historias HU-PM del Anexo A son propuestas para revisión y no sustituyen la aprobación del Product Owner.',
66:'Verificación de requisitos no funcionales: medir tiempos de respuesta, seguridad, precisión y disponibilidad en un entorno aislado, con carga y ventana de observación documentadas.',
67:'Numeración única: conservar los CP existentes y referenciarlos junto a su archivo. Reservar CP-PM-001 a CP-PM-012 para las propuestas de este plan; no renumerar automáticamente los casos existentes.',
81:'Captura de tres pesos, validación de discrepancia y cálculo de merma en ProductionLogsService.logTripleWeight y TripleWeightModal. La adquisición real desde báscula digital se tratará como integración pendiente de verificar con equipo y protocolo disponibles.',
90:'Registro de pesaje por triple factor y cálculo automático de merma; verificar por separado la captura manual y la integración física con báscula, sin dar esta última por acreditada.',
103:'Exclusión funcional de 2FA; no implica considerar bajo el riesgo global de autenticación. Mantener pruebas de contraseña, token, bloqueo y autorización.',
114:'Se limita la aceptación integral de RF-005 hasta aprobar su ficha. El código sí contiene órdenes: probar como dependencia la asignación, los cambios de estado y los datos usados por inventario, cronometría y pesaje.',
115:'Riesgo: flujos dependientes pueden aprobar parcialmente sin acreditar todos los requisitos de órdenes. Documentar el alcance de cada caso y no excluir las pruebas críticas de dependencia.',
122:'Jest 30 con ts-jest y Nest TestingModule. Ejecutar funciones/servicios con datos controlados, éxito, errores y límites. Un mock verifica el comportamiento aislado, no la conexión real a la base de datos.',
125:'Usar HTTP real, API y PostgreSQL de pruebas con datos reiniciables. Las carpetas llamadas integracion contienen casos con mocks y no acreditan por su nombre este nivel de integración.',
134:'Medir latencia y tasa de fallos en endpoints reales; revisar TLS, disponibilidad y precisión de pesajes. JWT representa un token: no cifra el canal de transporte.',
137:'Probar peticiones anónimas, cada rol, actor manipulado, cuenta inactiva, bloqueo de intentos, expiración/revocación y ausencia de password en respuestas. Relacionar la regresión con TEC-001 a TEC-004 del informe final.',
146:'La compilación y las pruebas unitarias están disponibles; para integración real se requiere además API operativa, BD aislada y datos cargados. No confundir las dos etapas.',
150:'Ejecutar y aprobar el 100 % de los casos de prioridad ALTA y todos los flujos críticos acordados; toda excepción requiere aceptación expresa del Product Owner.',
153:'La tasa de fallos debe ser menor al 10 % y no puede incluir fallos críticos o de alta prioridad abiertos. Todo defecto restante debe tener decisión, responsable y evidencia registrada.',
154:'Presentar indicadores con alcance, fecha, versión, casos ejecutados/aprobados/fallidos/bloqueados y cobertura; mantener separados los resultados de Jest y SonarQube.',
155:'Entregar informe final y evidencias reproducibles. El cierre global no está aprobado: existen TEC-001/TEC-002 críticos y validaciones pendientes.',
181:'Versión 1.2 actualizada con entorno técnico, trazabilidad propuesta y estado del corte de septiembre.',
184:'Banco ejecutable en test/unitarias, test/integracion, Skainet-API/pruebas y Skainet-API/src. Llevar los casos al formato P-QA-10 con paso, dato, resultado esperado y evidencia; el Excel aún requiere consolidación.',
187:'Sección 14 y Anexos A/B de este plan; consolidar la relación RF–HU–CP con la matriz oficial cuando se aprueben las HU.',
190:'informe-evidencias-jest.json, informe-evidencias-lint.json y Skainet-API/coverage; añadir evidencia HTTP, Cypress, RNF y SonarQube al ejecutarlos.',
193:'TEC-001 a TEC-008 del Informe_Final_Pruebas_Skainet_Completado.docx. Trasladar a P-QA-12 y registrar responsable, estado y prueba de cierre.',
196:'Resumen verificado del Anexo C y reporte HTML de cobertura. Dashboard consolidado del sprint pendiente de actualización.',
199:'Informe_Final_Pruebas_Skainet_Completado.docx: resultados, alcance, riesgos y aprobación condicionada.',
205:'Entorno local Windows disponible para compilación y Jest. Para HTTP/BD/carga: instancia separada de PostgreSQL y proceso o contenedor de API/worker. Dimensionar recursos y registrar CPU, RAM y versiones antes de medir rendimiento.',
207:'Equipo con Node y navegador para pruebas web; báscula y dispositivo móvil únicamente cuando se incorpore esa integración al ciclo. Registrar las características del equipo de medición.',
217:'Chrome y Firefox: registrar versión exacta al ejecutar la aceptación. No se certifica compatibilidad entre navegadores mediante el build.',
219:'PostgreSQL con Prisma 5; schema en Skainet-API/prisma/schema.prisma. DATABASE_URL exclusiva del entorno de pruebas; nunca reiniciar una base de desarrollo o producción para esta campaña.',
221:'Node.js v22.13.1 incluido en Skainet-Node/node.exe. API y worker NestJS 11 / TypeScript 5; frontend React 19 / Vite 8. App móvil Flutter/Dart presente, fuera de la ejecución Jest.',
243:'Jose Alejandro Gualguan Rivera',244:'Diseño y pruebas funcionales',246:'Jose Alejandro Gualguan Rivera',247:'Pruebas de API',249:'Jose Alejandro Gualguan Rivera',250:'Trazabilidad y métricas',252:'Jose Alejandro Gualguan Rivera',
255:'Preparar una sesión de repaso de Jest/mocks, lectura de LCOV, pruebas HTTP por rol, manejo seguro de datos y formato P-QA-10. Confirmar dominio de las herramientas antes de asignar ejecución independiente; no se presupone formación ya realizada.',
257:'Backend: API REST NestJS 11 y TypeScript 5 en Skainet-API; worker NestJS en Skainet-Backend. Puertos configurados por defecto: 3000 y 3001, respectivamente, modificables mediante PORT.',
258:'Frontend: React 19, JavaScript/JSX, React Router 7 y Vite 8 en Skainet-Frontend. Compilación verificada. Cypress contiene login.cy.js y navegacion.cy.js, aún sin ejecución acreditada.',
259:'Autenticación: cuentas ficticias independientes para Super Administrador, Administrador y Joyero; añadir cuentas inactivas y bloqueadas. El enum incluye además Lider de Taller y un alias DUENO; el equipo debe decidir su alcance. No reutilizar datos ni claves reales.',
260:'Datos: usar el catálogo y escenarios del Anexo B. Crear materiales con stock cero, suficiente e insuficiente; usuarios de cada rol; órdenes asignadas/no asignadas; registros de tiempo y tres pesos. Las pruebas Jest actuales usan mocks; la carga real en PostgreSQL sigue pendiente.',
261:'Documentación configurada: API http://localhost:3000/api y worker http://localhost:3001/api, una vez iniciados los servicios. Swagger está definido en cada src/main.ts. Son URLs locales configuradas, no enlaces de un despliegue verificado.',
269:'El ciclo organiza planificación, diseño, ejecución y evaluación tomando la serie ISO/IEC/IEEE 29119 como referencia de proceso y documentación. Registrar versión de código, datos, resultado esperado/obtenido y evidencia para repetir cada ejecución.',
283:'R / A',289:'R / A',295:'A',301:'A',307:'A',313:'A',319:'A',
330:'R = Responsable, A = Aprobador, C = Consultado, I = Informado. Proyecto individual: Jose Alejandro Gualguan Rivera asume desarrollo y QA. El instructor/Product Owner valida la aceptación de negocio y su nombre queda por confirmar. La autorrevisión técnica no equivale a una auditoría independiente ni a una firma de aprobación.',
339:'Jose Alejandro Gualguan Rivera',343:'Jose Alejandro Gualguan Rivera',347:'Jose Alejandro Gualguan Rivera',351:'Jose Alejandro Gualguan Rivera',355:'Jose Alejandro Gualguan Rivera',359:'Jose Alejandro Gualguan Rivera',363:'Jose Alejandro Gualguan Rivera',367:'Jose Alejandro Gualguan Rivera',368:'Revisión independiente propuesta',371:'Instructor por confirmar',375:'Jose Alejandro e instructor',
331:'13.3. Cronograma original del Sprint QA-02',
378:'Antes de iniciar cada campaña, QA debe confirmar disponibilidad del entorno, la BD aislada y la versión desplegada. El cronograma de julio es una línea base histórica, no un registro de ejecución.',
379:'Objetivo propuesto: atender defectos críticos dentro de 24 horas desde su clasificación. El líder de desarrollo debe confirmar capacidad y plazo; si no se cumple, se suspende la liberación afectada.',
382:'Trazabilidad al corte del 14/09/2026. Los CP indicados existen en archivos de test; su aprobación con mocks no prueba por sí sola el cumplimiento del requisito. HU-PM-001 a HU-PM-007 son propuestas del Anexo A pendientes de validación. Los casos CP-PM corresponden a nuevas pruebas diseñadas en el Anexo B, todavía no implementadas ni ejecutadas.',
391:'HU-PM-001 (propuesta)',392:'CP-001 a CP-003; CP-006 a CP-068. Archivos usuarios y usuarios-joyeros. Complementar CP-PM-001/002.',
396:'HU-PM-002 (propuesta)',397:'CP-004/005 y CP-069 a CP-078; autenticacion.*.spec.ts. Complementar CP-PM-003/004.',
401:'HU-PM-003 (propuesta)',402:'CP-085 a CP-094; inventario.*.spec.ts. Complementar CP-PM-005/006.',
405:'RF-005: módulo implementado, ficha formal pendiente de aprobación.',406:'HU-PM-004 (propuesta)',407:'CP-095 a CP-105; ordenes.*.spec.ts. Validar las dependencias usadas por los casos de producción.',
411:'HU-PM-005 (propuesta)',412:'CP-PM-007/008 propuestos; logPhaseTime sin cobertura directa en la ejecución de referencia.',
416:'HU-PM-006 (propuesta)',417:'CP-106 a CP-113 relacionados con producción/mermas; CP-PM-009/010 para logTripleWeight, aún no ejecutados.',
421:'HU-PM-007 (propuesta)',422:'CP-126 a CP-132; busqueda.*.spec.ts. Complementar CP-PM-011.',
424:'Clasificar los defectos por impacto y registrar ID, módulo, versión, pasos, datos ficticios, esperado/obtenido, evidencia, responsable y estado. Flujo: abierto → asignado → corregido → pendiente de revalidación → cerrado o reabierto. TEC-001 a TEC-008 del informe final son la lista inicial; no se cierran por la mera aprobación de la suite Jest.',
470:'Referencia documental: Requerimientos_Funcionales.docx citado en el informe final. El responsable documental debe confirmar su ubicación oficial; no se declara un enlace inexistente.',
472:'Anexo A de este plan: HU-PM-001 a HU-PM-007 propuestas. La formalización de las 22 HU observadas en el checklist continúa pendiente.',
474:'Antecedente de 22 CU recogido en el informe final. Confirmar el archivo oficial de casos extendidos antes de aprobar trazabilidad documental.',
476:'Banco ejecutable: test/unitarias/, test/integracion/, Skainet-API/pruebas/ y src/**/*.spec.ts. Casos nuevos: Anexo B. Excel P-QA-10 pendiente de consolidación.',
478:'Sección 14 de este plan, con IDs existentes/propuestos. Excel P-QA-11 pendiente de consolidación y aprobación.',
480:'Informe_Final_Pruebas_Skainet_Completado.docx, sección 8: TEC-001 a TEC-008. Registro P-QA-12 pendiente de seguimiento formal.',
482:'https://github.com/Skainet-Code/Skainet_code_Completo.git',
484:'API: http://localhost:3000/api. Worker: http://localhost:3001/api. Requieren servicios iniciados en el equipo local.',
490:'JWT: JSON Web Token. Formato de token usado en autenticación; la autorización requiere verificar permisos y JWT no sustituye el cifrado TLS.',
}
for i,t in changes.items():ps[i].text=t

old=next(t for t in doc.tables if t.cell(0,0).text=='Actividad')
raci=doc.add_table(rows=1,cols=3)
for c,v in zip(raci.rows[0].cells,['Actividad','Jose Alejandro Gualguan Rivera','Instructor / Product Owner']):c.text=v
for activity in ['Plan Maestro de Pruebas','Casos de Prueba','Matriz de Trazabilidad','Pruebas Funcionales','Pruebas de API','Registro de Defectos','Dashboard de Calidad','Informe Ejecutivo']:
 vals=[activity,'R' if activity=='Informe Ejecutivo' else 'R / A','A' if activity=='Informe Ejecutivo' else 'C']
 for c,v in zip(raci.add_row().cells,vals):c.text=v
old._tbl.addprevious(raci._tbl);old._tbl.getparent().remove(old._tbl)
raci.autofit=False
for col,w in zip(raci.columns,[2.1,2.3,2.1]):col.width=Inches(w)
for row in raci.rows:
 for c,w in zip(row.cells,[2.1,2.3,2.1]):c.width=Inches(w)

# Extend the existing tools table instead of leaving automation undocumented.
tools=next(t for t in doc.tables if t.cell(0,0).text=='Herramienta')
for vals in [('Jest / ts-jest','Pruebas de API con TypeScript; JSON de resultados y cobertura LCOV/HTML.'),('Cypress','Casos web de login y navegación existentes; ejecución E2E pendiente.'),('SonarQube / SonarScanner','Análisis estático e importación LCOV; servidor, credenciales y ejecución pendientes.'),('ESLint / Vite / TypeScript','Verificación de lint web y compilación de frontend, API y worker.')]:
 for c,v in zip(tools.add_row().cells,vals):c.text=v

def heading(t):
 p=doc.add_paragraph(t)
 p.paragraph_format.keep_with_next=True
 p.paragraph_format.space_before=Pt(14)
 p.paragraph_format.space_after=Pt(6)
 for r in p.runs:r.bold=True;r.font.size=Pt(14)
 return p
def para(t):
 p=doc.add_paragraph(t)
 if t.startswith(('Como ','Acción:','Desde la raíz','Para conservar')):p.paragraph_format.keep_with_next=True
 return p
def table(headers,rows,widths):
 t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for vals in rows:
  for c,v in zip(t.add_row().cells,vals):c.text=str(v)
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for row in t.rows:
  for c,w in zip(row.cells,widths):c.width=Inches(w)
 return t

doc.add_page_break()
heading('Anexo A Historias de usuario propuestas')
para('Estas siete historias resumen las épicas para permitir revisar la trazabilidad técnica. Requieren aprobación del Product Owner y descomposición en historias del backlog; no equivalen al cierre de las 22 HU no conformes del checklist. Estimación y sprint de implementación se acuerdan en planificación, sin asignar puntos ni compromisos no confirmados.')
stories=[
('HU-PM-001','Gestión de usuarios','Como responsable jerárquico, quiero administrar únicamente las cuentas permitidas por mi rol para mantener la segregación de funciones.','Dado un Administrador autenticado, cuando intenta modificar a otro Administrador o al Super Administrador, entonces se rechaza la operación sin cambiar los datos. Dado un actor que intenta desactivarse, entonces se rechaza la solicitud y se conserva su cuenta.'),
('HU-PM-002','Autenticación','Como usuario, quiero iniciar sesión con identificación y contraseña y gestionar mi clave para acceder con seguridad.','Dadas credenciales válidas de una cuenta activa, cuando inicio sesión, entonces recibo un token sin datos sensibles. Dada una cuenta inactiva o bloqueada, se rechaza el acceso. Verificar con el Product Owner el umbral propuesto de cinco fallos y quince minutos de bloqueo.'),
('HU-PM-003','Inventario','Como Administrador, quiero registrar entradas y salidas trazables para conocer el stock real de materiales.','Dado stock de 10 g, cuando registro una salida válida de 3 g vinculada a una OT, entonces el saldo es 7 g y existe un movimiento de kardex. Si el movimiento falla, ninguna de las dos escrituras se confirma.'),
('HU-PM-004','Órdenes de trabajo','Como Administrador, quiero asignar órdenes y consultar su estado para coordinar la producción del taller.','Dada una orden y un Joyero habilitado, cuando realizo una asignación válida, entonces ambos quedan relacionados. Una transición no permitida se rechaza y los procesos dependientes conservan trazabilidad; formalizar estados y reglas de RF-005 antes de aceptación integral.'),
('HU-PM-005','Cronometría','Como Joyero, quiero iniciar, pausar, reanudar y finalizar el tiempo de una fase para registrar mi trabajo.','Dada una OT asignada, cuando ejecuto una secuencia válida, entonces se registran actor, fase, acción y tiempo sin valores negativos. Una fase o transición inválida se rechaza; aprobar la matriz de transiciones antes de cerrar el caso.'),
('HU-PM-006','Pesaje triple','Como Joyero, quiero registrar tres pesos y conocer la merma para controlar el material utilizado.','Dados tres pesos positivos, cuando su diferencia está dentro de la tolerancia aprobada, entonces se registra el pesaje y la merma con precisión acordada. Si se excede la tolerancia, se rechaza y no se persiste. Resolver la diferencia entre 0,1 en el código y 0,05 g en el comentario.'),
('HU-PM-007','Búsqueda','Como usuario, quiero buscar información autorizada por texto y filtros para encontrar registros sin acceder a datos ajenos.','Dado un Joyero autenticado, cuando consulta una búsqueda global, entonces recibe únicamente información dentro de su alcance; ningún resultado incluye contraseñas. Sin coincidencias se presenta un resultado vacío comprensible.'),
]
for code,title,story,criteria in stories:
 heading(code+' '+title);para(story);para('Criterios propuestos: '+criteria)
para('Definition of Done propuesta: criterios aprobados, implementación revisada, pruebas pertinentes aprobadas, sin defectos críticos/altos abiertos, trazabilidad RF–HU–CP actualizada y evidencia archivada. Aprobación de negocio pendiente.')

heading('Anexo B Datos y casos complementarios')
para('Campaña propuesta para una base de datos aislada. Todos los identificadores y valores siguientes son ficticios. No se ha ejecutado esta carga ni estos casos CP-PM; los resultados esperados son objetivos de verificación, no resultados observados.')
table(['Conjunto','Datos de preparación'],[
('Usuarios','QA-SA-001 (Super Administrador), QA-AD-001/002 (Administradores), QA-JO-001/002 (Joyeros), más una cuenta inactiva y otra bloqueada. Generar claves de prueba privadas.'),
('Materiales','QA-MAT-001 oro 18k: stock 10,00 g; QA-MAT-002 plata: 0,00 g; QA-MAT-003 esmeralda: 2,00 ct; QA-MAT-004 diamante inactivo. Verificar catálogo de unidades aprobado.'),
('Órdenes','QA-OT-001 asignada a QA-JO-001; QA-OT-002 asignada a QA-JO-002; una OT cerrada y un identificador inexistente.'),
('Pesos y tiempos','Pesos 10,00/10,02/10,01; 10,00/10,00/10,10; 10,00/10,00/10,11; cero, negativos y no numéricos. Duraciones 0, 60 y -1 s. Probar límites con el umbral aprobado.')],[1.3,5.2])
cases=[
('CP-PM-001','Alta · TEC-001','Consultar GET /users/:id sin token y con rol ajeno.','La consulta no autorizada se rechaza y ninguna respuesta autorizada contiene password.'),
('CP-PM-002','Alta · TEC-004','Intentar autodesactivación y modificación de Super Administrador como Administrador; manipular actorId en el cuerpo.','Rechazo sin cambios; identidad tomada del usuario autenticado y auditoría consistente.'),
('CP-PM-003','Alta · TEC-003','Realizar cinco fallos de login usando el UsersService real; luego intentar clave correcta durante el bloqueo.','Con la política 5/15 min aprobada, acceso bloqueado, contador y auditoría persistidos; comprobar desbloqueo al finalizar el plazo.'),
('CP-PM-004','Alta · TEC-002/004','Intentar iniciar la API sin JWT_SECRET, usar token alterado/caducado y token de una cuenta desactivada.','Configuración insegura rechazada al iniciar; tokens alterados/caducados y cuenta inactiva sin acceso.'),
('CP-PM-005','Alta · TEC-005','Con 10 g de stock, solicitar en paralelo dos salidas de 7 g sobre la misma referencia.','Como máximo una salida confirmada; saldo coherente, nunca negativo; kardex coincide con las operaciones aceptadas.'),
('CP-PM-006','Alta · TEC-005','Provocar fallo controlado al crear kardex en una entrada/salida.','Rollback completo: el stock no cambia si el movimiento no se registra.'),
('CP-PM-007','Alta · RF-006','POST /production-logs/timer: iniciar, pausar, reanudar y finalizar una fase de OT asignada.','Secuencia, actor y duración consistentes; consulta recupera los registros de esa OT.'),
('CP-PM-008','Alta · RF-006','Enviar fase inválida, acción desconocida, duración negativa y OT ajena/inexistente.','Cada entrada inválida o acceso ajeno se rechaza sin persistencia; comportamiento aprobado y documentado.'),
('CP-PM-009','Alta · RF-007','POST /production-logs/weights con 10,00/10,02/10,01; consultar el pesaje.','Registro con precisión acordada, actor/OT correctos y fórmula de merma verificada, incluyendo tratamiento de ganancia de peso según regla aprobada.'),
('CP-PM-010','Alta · TEC-008','Probar diferencia exactamente igual al umbral y superior a él; cero, negativos y no numéricos.','Valores inválidos rechazados; frontera aprobada aceptada y exceso rechazado, sin registros parciales.'),
('CP-PM-011','Alta · RF-008','GET /search como Joyero buscando órdenes y usuarios fuera de su alcance.','Ningún dato ajeno o sensible; filtros y resultado vacío correctos.'),
('CP-PM-012','Media · RNF-001','Medir una campaña propuesta de 50 usuarios virtuales durante 5 min tras calentamiento de 1 min en endpoints reales.','Registrar máximo, p50, p95, p99 y errores; contrastar cada respuesta con el límite de 3 s del RNF. Carga y perfil requieren aprobación; no usar duraciones simuladas.')
]
for code,priority,steps,expected in cases:
 heading(code+' '+priority.replace(' · ',' '));para('Acción: '+steps);para('Resultado esperado: '+expected)
para('Preparación y limpieza: apuntar DATABASE_URL a una instancia identificada como QA; aplicar migraciones verificadas, crear los datos ficticios y registrar sus IDs. Al terminar, limpiar únicamente la instancia QA mediante el procedimiento aprobado. No ejecutar reset, seed o borrado contra desarrollo/producción. El responsable registra versión del esquema y evidencia de restauración.')

heading('Anexo C Ejecución reproducible y seguimiento')
para('Evidencia de referencia del 14/09/2026: 239/239 pruebas, 58/58 suites; 473/902 líneas cubiertas. Sentencias 51,16 %, ramas 53,73 %, funciones 50,93 % y líneas 52,43 %. API y worker compilan; frontend Vite compila; ESLint: 0 errores y 3 advertencias. Son resultados de la ejecución local anterior, no de los casos CP-PM propuestos ni del servidor SonarQube.')
para('Desde la raíz Skainet_FI, ejecutar en PowerShell:')
para('.\\ejecutar_cobertura.bat')
para('Para conservar el detalle de Jest:')
para('.\\Skainet-Node\\node.exe Skainet-API/node_modules/jest/bin/jest.js --config Skainet-API/jest.config.json --runInBand --json --outputFile=informe-evidencias-jest.json')
para('Consultar Skainet-API/coverage/lcov-report/index.html e informe-evidencias-jest.json. La configuración de Jest de la API incluye src, pruebas y las carpetas test de la raíz; no sumar copias fuera de esa selección. Corregir por separado las rutas backend/ del Jest del worker antes de acreditar sus pruebas.')
para('SonarQube: generar primero LCOV; configurar URL, clave del proyecto y SONAR_TOKEN de forma privada; ejecutar SonarScanner desde la raíz con sonar-project.properties. Conservar el enlace del análisis procesado y su Quality Gate. No trasladar porcentajes de Jest como si fueran resultados de SonarQube, ni declarar deuda técnica o vulnerabilidades no medidas.')
table(['Actividad de cierre propuesta','Horas','Responsable por rol','Evidencia de salida'],[
('Preparación de entorno, datos y criterios','4 h','QA Lead + Desarrollo','Entorno QA identificado y criterios acordados.'),
('Regresión de seguridad y permisos','10 h','QA + Desarrollo','TEC-001 a TEC-004 corregidos y revalidados.'),
('Inventario, tiempos y pesaje','10 h','QA + Desarrollo','CP-PM-005 a CP-PM-010 con evidencia.'),
('Web, búsqueda y dependencias','6 h','QA','Flujos por rol y búsqueda verificados.'),
('RNF y SonarQube','6 h','QA + Desarrollo','Métricas reales y análisis procesado.'),
('Trazabilidad, informe y revisión','4 h','QA Lead + Product Owner','Artefactos consistentes y decisión de aceptación.')],[2.35,.55,1.45,2.15])
para('Total propuesto: 40 horas de actividad QA, sujeto a acuerdo de inicio y disponibilidad. No incluye una estimación de desarrollo para resolver defectos ni acredita el 99 % de disponibilidad: esa medición requiere una ventana de observación acordada y monitoreo continuo. Las fechas del nuevo ciclo se fijan al aprobar su inicio.')
para('Indicadores: aprobación = casos aprobados / casos ejecutados × 100; fallos = casos fallidos / casos ejecutados × 100; registrar bloqueados y no ejecutados por separado. Cobertura de código corresponde al reporte instrumentado y no a cobertura total de requisitos. Disponibilidad = tiempo disponible / tiempo de observación × 100, con exclusiones aprobadas y evidencia de monitoreo.')
para('Estado de cierre: condicionado. TEC-001/002 críticos impiden recomendar liberación con datos reales. No hay firmas de aprobación acreditadas, ni resultados de integración real, Cypress, Flutter, carga o disponibilidad incluidos en las 239 pruebas de referencia.')
heading('Referencias de consulta')
para('ISO/IEC/IEEE 29119-1:2022, conceptos de pruebas: https://www.iso.org/standard/81291.html')
para('ISO/IEC/IEEE 29119-2:2021, procesos de pruebas: https://www.iso.org/standard/79428.html')
para('ISO/IEC 25010:2023, modelo de calidad de producto: https://www.iso.org/standard/78176.html')
para('Fuentes del proyecto: manifiestos package.json, esquema Prisma, configuración Jest, casos bajo test/, reporte JSON de resultados y el informe final de pruebas completado.')

# Format original template and extensions with consistent tables and headings.
for p in doc.paragraphs:
 if re.match(r'^(\d+\.\s|\d+\.\d+\.|Anexo [ABC]|Referencias)',p.text):
  p.paragraph_format.keep_with_next=True
  p.paragraph_format.keep_together=True
 for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
for t in doc.tables:
 pr=t._tbl.tblPr
 borders=pr.find(qn('w:tblBorders'))
 if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
 for edge in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr()
  if not trpr.xpath('./w:cantSplit'):trpr.append(OxmlElement('w:cantSplit'))
  for h in trpr.xpath('./w:trHeight'):trpr.remove(h)
  if ri==0 and len(t.rows)>2:
   if not trpr.xpath('./w:tblHeader'):trpr.append(OxmlElement('w:tblHeader'))
  for c in row.cells:
   cp=c._tc.get_or_add_tcPr()
   margins=cp.find(qn('w:tcMar'))
   if margins is None:margins=OxmlElement('w:tcMar');cp.append(margins)
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');margins.append(el)
   for p in c.paragraphs:
    p.paragraph_format.keep_with_next=False
    for r in p.runs:
     r.font.color.rgb=RGBColor(0,0,0)
     if len(t.columns)>=4:r.font.size=Pt(9)
   # Preserve metadata tables without artificial header treatment.
   if not (len(t.columns)==2 and t.cell(0,0).text in ['Código','Empresa / Entidad']):
    sh=cp.find(qn('w:shd'))
    if sh is None:sh=OxmlElement('w:shd');cp.append(sh)
    sh.set(qn('w:fill'),'203864' if ri==0 else ('F2F4F7' if ri%2==0 else 'FFFFFF'))
    if ri==0:
     for p in c.paragraphs:
      for r in p.runs:r.font.color.rgb=RGBColor(255,255,255);r.bold=True
for p in doc.element.body.xpath('.//w:p'):
 text=''.join(p.xpath('.//w:t/text()'))
 assert '[Completar' not in text and '[Nombre]' not in text,text
for r in doc.paragraphs[0].runs:r.font.color.rgb=RGBColor(0,0,0)
doc.core_properties.title='Plan maestro de pruebas de Skainet'
doc.core_properties.subject='Versión 1.2 actualizada al 14 de septiembre de 2026'
out=base/'Plan_Maestro_de_Pruebas_Skainet_Completado.docx'
doc.save(out)
print(out)
